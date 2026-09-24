"""Excel workbook parsing.

Pure parsing of the Excel (.xlsx) export. Works on a filesystem path, raw
bytes (from an HTTP upload), or any file-like object -- the service receives
uploads and parses them in memory, never touching disk.
"""
from __future__ import annotations

import io
import warnings
from dataclasses import dataclass, field
from typing import Any, BinaryIO

warnings.filterwarnings("ignore")  # silence openpyxl data-validation warning
import openpyxl


@dataclass
class ParsedWorkbook:
    """Everything the mapper needs, extracted from one workbook."""

    info: dict[str, Any]                       # citation: title/authors/contacts/description/subject
    field_dict: dict[str, dict[str, dict]]     # template -> {display label -> {typeName,type,required}}
    meta_sheets: dict[str, list[dict]]         # sheet title -> [ {display label: value}, ... ]
    sheet_names: list[str] = field(default_factory=list)


def _open(source) -> "openpyxl.Workbook":
    """Load a workbook from a path, bytes, or file-like object."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    return openpyxl.load_workbook(source, data_only=True, read_only=True)


def _clean(s) -> str | None:
    """Normalize a cell: NBSP -> space, collapse runs, trim. None stays None."""
    if s is None:
        return None
    t = " ".join(str(s).replace("\xa0", " ").split())
    return t or None


def read_template_info(wb) -> dict[str, Any]:
    """Citation metadata from the TemplateInfo sheet -> a simple dict.

    Layout (label/value pairs, indentation groups sub-rows under a header):

        Account Information          <- optional; also matches "Depositor Account
          Username / Email              Information". Identifies the depositing account.
        Title
        Author        -> Name / Affiliation
        Point of Contact -> Name / E-mail
        Description
        Subject
        Depositor                    <- optional

    Repeatable fields (Author, Point of Contact, Subject) may appear more than
    once and all values are kept. Two ways to add another entry work equally:
      * repeat the header block (another `Author` / `Point of Contact`), or
      * repeat the sub-rows under one header (a second `Name`/`E-mail` starts a
        new entry once the current one already has that sub-field).

    `Email` under Account Information and `E-mail` under Point of Contact are
    disambiguated by which section we are currently inside.
    """
    ws = wb["TemplateInfo"]
    info: dict[str, Any] = {"authors": [], "contacts": [], "subjects": [], "account": {}}
    cur_author = cur_contact = None
    section: str | None = None          # "account" | "author" | "contact" | None

    def entry_for(entries, cur, fld):
        """Write target for `fld`: reuse the current entry, or start a new one
        if there is none yet or it already holds this sub-field."""
        if cur is None or fld in cur:
            cur = {}
            entries.append(cur)
        return cur

    def _truthy(x) -> bool:
        return str(x).strip().lower() in ("true", "1", "yes", "y", "x")

    for row in ws.iter_rows(values_only=True):
        label = row[0] if len(row) > 0 else None
        value = row[1] if len(row) > 1 else None
        flag = row[2] if len(row) > 2 else None      # column C: "Contributor Role" true/false
        if label is None:
            continue
        key = (_clean(label) or "").lower()
        val = _clean(value)

        if "account information" in key:      # "Account Information" or "Depositor Account Information"
            section, cur_author, cur_contact = "account", None, None
        elif key == "author":
            section, cur_author, cur_contact = "author", None, None
        elif key == "point of contact":
            section, cur_contact, cur_author = "contact", None, None
        # --- top-level fields: they also close any open section ---
        elif key == "title":
            info["title"] = val
            section, cur_author, cur_contact = None, None, None
        elif key == "description":
            info["description"] = val
            section, cur_author, cur_contact = None, None, None
        elif key == "subject":
            if val:
                info["subjects"].append(val)
            section, cur_author, cur_contact = None, None, None
        elif key == "depositor":
            info["depositor"] = val
            section, cur_author, cur_contact = None, None, None
        # --- sub-rows, resolved against the current section ---
        elif key == "username":
            if section == "account":
                info["account"]["username"] = val
            elif section == "author" and cur_author is not None:
                cur_author["username"] = val
        elif key in ("e-mail", "email"):
            if section == "account":
                info["account"]["email"] = val
            elif section == "contact":
                cur_contact = entry_for(info["contacts"], cur_contact, "email")
                cur_contact["email"] = val
            elif section == "author" and cur_author is not None:
                cur_author["email"] = val
        elif key == "name":
            if section == "contact":
                cur_contact = entry_for(info["contacts"], cur_contact, "name")
                cur_contact["name"] = val
            elif section == "author":
                cur_author = entry_for(info["authors"], cur_author, "name")
                cur_author["name"] = val
        elif key == "affiliation":
            if section == "author":
                cur_author = entry_for(info["authors"], cur_author, "affiliation")
                cur_author["affiliation"] = val

        # Column C on any author row carries that author's "Contributor Role" flag.
        if section == "author" and cur_author is not None and flag is not None \
                and str(flag).strip() != "":
            cur_author["contributor"] = _truthy(flag)

    info["authors"] = [a for a in info["authors"] if a.get("name")]
    info["contacts"] = [c for c in info["contacts"] if c.get("email") or c.get("name")]
    info["account"] = {k: v for k, v in info["account"].items() if v}
    if info["subjects"]:
        info["subject"] = info["subjects"][0]  # scalar kept for the required-field check
    return info


def read_field_dictionary(wb) -> dict[str, dict[str, dict]]:
    """template -> {display label -> {typeName, type, required}}."""
    ws = wb["Field Dictionary"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    needed = ("Template", "Field Name (ID)", "Display Label", "Type", "Required")
    missing = [c for c in needed if c not in header]
    if missing:
        raise ValueError(f"Field Dictionary is missing columns: {missing}")
    idx = {n: header.index(n) for n in needed}
    out: dict[str, dict[str, dict]] = {}
    for r in rows[1:]:
        if not r or r[idx["Template"]] is None:
            continue
        tmpl = str(r[idx["Template"]]).strip()
        label = str(r[idx["Display Label"]]).strip()
        out.setdefault(tmpl, {})[label] = {
            "typeName": str(r[idx["Field Name (ID)"]]).strip(),
            "type": str(r[idx["Type"]]).strip().lower(),
            "required": str(r[idx["Required"]]).strip().lower() == "yes",
        }
    return out


def read_metadata_sheets(wb) -> dict[str, list[dict]]:
    """Every 'Metadata_*' sheet -> list of dict rows keyed by display label.

    Empty cells are skipped, so only fields actually filled in are emitted.
    """
    out: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        if not ws.title.startswith("Metadata_"):
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        entries = []
        for r in rows[1:]:
            if r is None or all(c is None for c in r):
                continue
            entry = {h: c for h, c in zip(headers, r) if h and c is not None and str(c).strip() != ""}
            if entry:
                entries.append(entry)
        out[ws.title] = entries
    return out


def match_template(sheet_title: str, field_dict: dict) -> str | None:
    """Excel truncates sheet names to 31 chars; match back to the full template."""
    stub = sheet_title[len("Metadata_"):].strip()
    for tmpl in field_dict:
        if tmpl.startswith(stub) or stub.startswith(tmpl[: len(stub)]):
            return tmpl
    return None


def resolve_template(sheet_title: str, field_dict: dict,
                     sheet_map: dict[str, str] | None = None) -> str | None:
    """Which template a sheet maps to. An explicit `sheet_map` (sheet title ->
    template, from the user's manual matching) wins; an empty/None value there
    means "skip this sheet". Otherwise fall back to name-based auto-matching."""
    if sheet_map is not None and sheet_title in sheet_map:
        chosen = (sheet_map.get(sheet_title) or "").strip()
        return chosen or None
    return match_template(sheet_title, field_dict)


def parse_workbook(source) -> ParsedWorkbook:
    """Parse an Excel .xlsx from a path, bytes, or file-like object."""
    wb = _open(source)
    try:
        return ParsedWorkbook(
            info=read_template_info(wb),
            field_dict=read_field_dictionary(wb),
            meta_sheets=read_metadata_sheets(wb),
            sheet_names=list(wb.sheetnames),
        )
    finally:
        wb.close()
