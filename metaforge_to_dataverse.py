#!/usr/bin/env python3
"""
metaforge_to_dataverse.py
-------------------------
Turn a MetaForge multi-metadata Excel export into a Dataverse
"create dataset" JSON payload, and (optionally) POST it to a Dataverse
collection.

Reference: https://guides.dataverse.org/en/6.6/api/native-api.html#create-dataset-command

The Excel is expected to have (as produced by MetaForge):
  * TemplateInfo                 -> citation metadata (Title/Author/Contact/Description/Subject)
  * Field Dictionary             -> per-field typeName, type (enum/text/int/url), required
  * Metadata_<block>  (>=1 row)  -> the actual values for each custom block
                                    (multiple rows = repeatable compound entries)

Two things are resolved from the *running* Dataverse when --server is given,
because the Excel alone cannot know them:
  * the internal block name  (JSON key, e.g. "crc1607SummaryOfSamples")
  * the compound-parent typeName + each field's typeClass / multiple flags
If --server is not given, the script infers them (shared field-id prefix for the
compound parent, and a slugified display name for the block key) and clearly
marks those as assumptions in the output.

Usage
-----
  # 1. Just generate the JSON (offline, uses inference):
  python metaforge_to_dataverse.py MetaForge_multi_2026-04-22.xlsx -o dataset.json

  # 2. Generate using the live block schema (recommended - exact typeNames):
  python metaforge_to_dataverse.py MetaForge_multi_2026-04-22.xlsx \
      --server https://134.95.195.250 --parent crc1218_testing \
      --token $API_TOKEN --insecure -o dataset.json

  # 3. Generate AND create the dataset:
  python metaforge_to_dataverse.py MetaForge_multi_2026-04-22.xlsx \
      --server https://134.95.195.250 --parent crc1218_testing \
      --token $API_TOKEN --insecure --post
"""
from __future__ import annotations
import argparse, json, re, sys, warnings
from pathlib import Path

warnings.filterwarnings("ignore")  # silence openpyxl data-validation warning
import openpyxl

try:
    import requests
except ImportError:
    requests = None  # only needed for --server / --post


# --------------------------------------------------------------------------- #
# Excel reading
# --------------------------------------------------------------------------- #
def load_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    return rows


def read_template_info(wb):
    """Citation metadata from the TemplateInfo sheet -> a simple dict."""
    ws = wb["TemplateInfo"]
    info = {"authors": [], "contacts": []}
    cur_author = cur_contact = None
    for label, value in ws.iter_rows(values_only=True):
        if label is None:
            continue
        key = str(label).strip()
        val = None if value is None else str(value).strip()
        low = key.lower()
        if low == "title":
            info["title"] = val
        elif low == "author":
            cur_author = {}
            info["authors"].append(cur_author)
            cur_contact = None
        elif low == "point of contact":
            cur_contact = {}
            info["contacts"].append(cur_contact)
            cur_author = None
        elif low == "name":
            if cur_contact is not None:
                cur_contact["name"] = val
            elif cur_author is not None:
                cur_author["name"] = val
        elif low == "affiliation":
            if cur_author is not None:
                cur_author["affiliation"] = val
        elif low == "e-mail" or low == "email":
            if cur_contact is not None:
                cur_contact["email"] = val
        elif low == "description":
            info["description"] = val
        elif low == "subject":
            info["subject"] = val
    # drop empties
    info["authors"] = [a for a in info["authors"] if a.get("name")]
    info["contacts"] = [c for c in info["contacts"] if c.get("email") or c.get("name")]
    return info


def read_field_dictionary(wb):
    """
    -> { template_display_name: {display_label: {"typeName","type","required"}} }
    'type' is the MetaForge type: enum / text / integer / url.
    """
    ws = wb["Field Dictionary"]
    rows = load_rows(ws)
    header = [str(h).strip() if h else "" for h in rows[0]]
    idx = {name: header.index(name) for name in
           ("Template", "Field Name (ID)", "Display Label", "Type", "Required")}
    out = {}
    for r in rows[1:]:
        if not r or r[idx["Template"]] is None:
            continue
        tmpl = str(r[idx["Template"]]).strip()
        label = str(r[idx["Display Label"]]).strip()
        out.setdefault(tmpl, {})[label] = {
            "typeName": str(r[idx["Field Name (ID)"]]).strip(),
            "type": str(r[idx["Type"]]).strip().lower(),
            "required": str(r[idx["Required"]]).strip().lower() == "yes",
        }
    return out


def read_metadata_sheets(wb):
    """
    Every 'Metadata_*' sheet -> list of dict rows keyed by display label.
    Returns { sheet_title: [ {label: value, ...}, ... ] }.
    Empty cells are skipped.
    """
    out = {}
    for ws in wb.worksheets:
        if not ws.title.startswith("Metadata_"):
            continue
        rows = load_rows(ws)
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        entries = []
        for r in rows[1:]:
            if r is None or all(c is None for c in r):
                continue
            entry = {}
            for h, c in zip(headers, r):
                if h and c is not None and str(c).strip() != "":
                    entry[h] = c
            if entry:
                entries.append(entry)
        out[ws.title] = entries
    return out


def match_template(sheet_title, field_dict):
    """
    Metadata sheet titles are truncated (Excel 31-char limit), e.g.
    'Metadata_CRC1607 Summary of Sam'. Match to the full template display name.
    """
    stub = sheet_title[len("Metadata_"):].strip()
    for tmpl in field_dict:
        if tmpl.startswith(stub) or stub.startswith(tmpl[:len(stub)]):
            return tmpl
    return None


# --------------------------------------------------------------------------- #
# Dataverse field-object builders
# --------------------------------------------------------------------------- #
def prim(type_name, value, multiple=False):
    return {"typeName": type_name, "multiple": multiple,
            "typeClass": "primitive", "value": value}


def cv(type_name, value, multiple=True):
    return {"typeName": type_name, "multiple": multiple,
            "typeClass": "controlledVocabulary", "value": value}


def compound(type_name, entries, multiple=True):
    return {"typeName": type_name, "multiple": multiple,
            "typeClass": "compound", "value": entries}


def build_citation(info):
    fields = []
    if info.get("title"):
        fields.append(prim("title", info["title"]))

    if info.get("authors"):
        authors = []
        for a in info["authors"]:
            obj = {"authorName": prim("authorName", a["name"])}
            if a.get("affiliation"):
                obj["authorAffiliation"] = prim("authorAffiliation", a["affiliation"])
            authors.append(obj)
        fields.append(compound("author", authors))

    if info.get("contacts"):
        contacts = []
        for c in info["contacts"]:
            obj = {}
            if c.get("name"):
                obj["datasetContactName"] = prim("datasetContactName", c["name"])
            if c.get("email"):
                obj["datasetContactEmail"] = prim("datasetContactEmail", c["email"])
            contacts.append(obj)
        fields.append(compound("datasetContact", contacts))

    desc = info.get("description")
    fields.append(compound("dsDescription",
        [{"dsDescriptionValue": prim("dsDescriptionValue", desc or "")}]))

    if info.get("subject"):
        fields.append(cv("subject", [info["subject"]]))

    return {"fields": fields, "displayName": "Citation Metadata"}


# --------------------------------------------------------------------------- #
# Server introspection (optional, recommended)
# --------------------------------------------------------------------------- #
def fetch_block_schema(server, parent, token, verify):
    """
    Returns { template_display_name: {
        "blockName": <internal name>,
        "fields": { typeName: {"typeClass","multiple","parent": <parent typeName or None>} }
    } } for every metadata block on the parent collection.
    """
    url = f"{server.rstrip('/')}/api/dataverses/{parent}/metadatablocks"
    params = {"returnDatasetFieldTypes": "true"}
    headers = {"X-Dataverse-key": token} if token else {}
    resp = requests.get(url, params=params, headers=headers, verify=verify, timeout=30)
    resp.raise_for_status()
    data = resp.json()["data"]
    schema = {}
    for block in data:
        display = block.get("displayName") or block.get("name")
        fields = {}

        def walk(field_map, parent_type=None):
            for tname, fdef in field_map.items():
                fields[tname] = {
                    "typeClass": fdef.get("typeClass"),
                    "multiple": fdef.get("multiple", False),
                    "parent": parent_type,
                }
                if fdef.get("childFields"):
                    walk(fdef["childFields"], tname)

        walk(block.get("fields", {}))
        schema[display] = {"blockName": block["name"], "fields": fields}
    return schema


def slug_block_name(display):
    """Fallback internal-name guess when we can't read it from the server."""
    parts = re.sub(r"[^0-9a-zA-Z ]+", " ", display).split()
    if not parts:
        return display
    return parts[0].lower() + "".join(p.capitalize() for p in parts[1:])


# --------------------------------------------------------------------------- #
# Build a custom (CRC1607) block
# --------------------------------------------------------------------------- #
def build_custom_block(template, entries, field_dict, schema):
    """
    entries : list of {display_label: value} (one per row)
    Produces a metadataBlocks[blockName] object.
    Returns (block_name, block_obj, assumptions[list of str]).
    """
    assumptions = []
    labels = field_dict[template]  # {display_label: {typeName,type,required}}
    # label -> typeName
    label_to_type = {lbl: meta["typeName"] for lbl, meta in labels.items()}

    # --- figure out block name + compound parent + typeClass/multiple ------- #
    if schema and template in schema:
        block_name = schema[template]["blockName"]
        sfields = schema[template]["fields"]

        def type_class(tn):
            return sfields.get(tn, {}).get("typeClass", "primitive")

        def parent_of(tn):
            return sfields.get(tn, {}).get("parent")

        # the compound parent is the shared parent of the child fields
        parents = {parent_of(tn) for tn in label_to_type.values() if parent_of(tn)}
        parent_type = next(iter(parents)) if len(parents) == 1 else None
        parent_multiple = sfields.get(parent_type, {}).get("multiple", True) if parent_type else True
    else:
        # ---- inference (no server) ----
        block_name = slug_block_name(template)
        assumptions.append(
            f"block internal name for '{template}' inferred as '{block_name}' "
            f"(confirm via /api/dataverses/<alias>/metadatablocks)")
        # compound parent = longest common prefix of the child field ids, trimmed at '_'
        tns = list(label_to_type.values())
        prefix = tns[0]
        for tn in tns[1:]:
            while not tn.startswith(prefix):
                prefix = prefix.rsplit("_", 1)[0] if "_" in prefix else ""
        parent_type = prefix.rstrip("_") or None
        if parent_type:
            assumptions.append(
                f"compound parent typeName inferred as '{parent_type}'")
        parent_multiple = True

        def type_class(tn):
            # map MetaForge type -> Dataverse typeClass via field dict
            for lbl, meta in labels.items():
                if meta["typeName"] == tn:
                    return "controlledVocabulary" if meta["type"] == "enum" else "primitive"
            return "primitive"

    # --- build one compound object per row ---------------------------------- #
    compound_entries = []
    for row in entries:
        obj = {}
        for label, value in row.items():
            tn = label_to_type.get(label)
            if not tn:
                continue  # column not in the dictionary; skip
            tc = type_class(tn)
            if tc == "controlledVocabulary":
                obj[tn] = cv(tn, str(value), multiple=False)
            else:
                # integers stay numeric-as-string per Dataverse convention
                obj[tn] = prim(tn, str(value))
        if obj:
            compound_entries.append(obj)

    if parent_type:
        block_obj = {
            "fields": [compound(parent_type, compound_entries, multiple=parent_multiple)],
            "displayName": template,
        }
    else:
        # no compound parent -> emit fields at top level from the FIRST row only
        assumptions.append(
            f"'{template}': no compound parent found; only the first row is emitted "
            f"as flat fields")
        flat = []
        if compound_entries:
            flat = list(compound_entries[0].values())
        block_obj = {"fields": flat, "displayName": template}

    return block_name, block_obj, assumptions


# --------------------------------------------------------------------------- #
# Main assembly
# --------------------------------------------------------------------------- #
def build_payload(xlsx_path, schema=None, description_override=None):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    info = read_template_info(wb)
    if description_override is not None:
        info["description"] = description_override
    field_dict = read_field_dictionary(wb)
    meta_sheets = read_metadata_sheets(wb)

    blocks = {"citation": build_citation(info)}
    assumptions, warnings_ = [], []

    # required citation checks
    for req_key, human in (("title", "Title"), ("authors", "Author Name"),
                           ("contacts", "Point of Contact Email"),
                           ("subject", "Subject")):
        if not info.get(req_key):
            warnings_.append(f"REQUIRED citation field missing: {human}")
    if not info.get("description"):
        warnings_.append("REQUIRED citation field missing: Description Text "
                         "(pass --description to supply one)")

    for sheet_title, entries in meta_sheets.items():
        template = match_template(sheet_title, field_dict)
        if not template:
            warnings_.append(f"Could not match sheet '{sheet_title}' to a template")
            continue
        if not entries:
            continue
        bname, bobj, ass = build_custom_block(template, entries, field_dict, schema)
        blocks[bname] = bobj
        assumptions.extend(ass)

    payload = {"datasetVersion": {"metadataBlocks": blocks}}
    return payload, assumptions, warnings_


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx", help="MetaForge Excel file")
    p.add_argument("-o", "--out", default="dataset.json", help="output JSON path")
    p.add_argument("--server", help="Dataverse base URL, e.g. https://134.95.195.250")
    p.add_argument("--parent", help="parent collection alias, e.g. crc1218_testing")
    p.add_argument("--token", help="Dataverse API token (X-Dataverse-key)")
    p.add_argument("--insecure", action="store_true",
                   help="skip TLS verification (self-signed local cert)")
    p.add_argument("--description", help="supply/override the dataset Description")
    p.add_argument("--post", action="store_true",
                   help="create the dataset after building the JSON")
    args = p.parse_args()

    schema = None
    if args.server and args.parent:
        if requests is None:
            sys.exit("`requests` is required for --server (pip install requests)")
        try:
            schema = fetch_block_schema(args.server, args.parent, args.token,
                                        verify=not args.insecure)
            print(f"[i] Read block schema from server "
                  f"({len(schema)} blocks): {', '.join(schema)}")
        except Exception as e:
            print(f"[!] Could not read block schema ({e}); falling back to inference.")

    payload, assumptions, warns = build_payload(
        args.xlsx, schema=schema, description_override=args.description)

    Path(args.out).write_text(json.dumps(payload, indent=2, ensure_ascii=False))
    print(f"[i] Wrote {args.out}")

    for w in warns:
        print(f"[WARN] {w}")
    for a in assumptions:
        print(f"[assumption] {a}")

    if args.post:
        if requests is None:
            sys.exit("`requests` is required for --post")
        if not (args.server and args.parent and args.token):
            sys.exit("--post needs --server, --parent and --token")
        url = f"{args.server.rstrip('/')}/api/dataverses/{args.parent}/datasets"
        r = requests.post(url, headers={"X-Dataverse-key": args.token,
                                        "Content-type": "application/json"},
                          data=json.dumps(payload), verify=not args.insecure, timeout=60)
        print(f"[i] POST {url} -> {r.status_code}")
        try:
            print(json.dumps(r.json(), indent=2))
        except Exception:
            print(r.text)
        r.raise_for_status()


if __name__ == "__main__":
    main()
